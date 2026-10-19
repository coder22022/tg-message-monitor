import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from group_export import (
    GROUP_CSV_HEADERS,
    MEMBER_CSV_HEADERS,
    _sort_dialogs_by_group_title,
    _sort_members_by_display_name,
    _write_export_workbook,
)


class WriteExportWorkbookTests(unittest.TestCase):
    def test_writes_groups_and_members_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            excel_path = Path(tmp) / "groups_export.xlsx"
            group_rows = [
                {
                    "group_id": 1,
                    "title": "G1",
                    "username": "g1",
                    "type": "supergroup",
                    "description": "d",
                    "member_count": 1,
                    "includes_all_members": True,
                    "online_count": 0,
                    "unread_count": 0,
                    "is_verified": False,
                    "is_scam": False,
                    "is_fake": False,
                    "avatar_path": "",
                    "exported_at": "2026-01-01",
                }
            ]
            member_rows = [
                {
                    "group_id": 1,
                    "group_title": "G1",
                    "user_id": 9,
                    "username": "u9",
                    "first_name": "A",
                    "last_name": "",
                    "display_name": "A",
                    "is_bot": False,
                    "is_premium": False,
                    "is_admin": False,
                    "is_creator": False,
                    "is_deleted": False,
                    "phone": "",
                    "avatar_path": "",
                }
            ]
            _write_export_workbook(excel_path, group_rows, member_rows)
            self.assertTrue(excel_path.is_file())

            workbook = load_workbook(excel_path, read_only=True)
            self.assertEqual(workbook.sheetnames, ["groups", "members"])
            groups = workbook["groups"]
            members = workbook["members"]
            self.assertEqual(
                [cell.value for cell in groups[1]],
                GROUP_CSV_HEADERS,
            )
            self.assertEqual(groups[2][0].value, 1)
            self.assertEqual(
                [cell.value for cell in members[1]],
                MEMBER_CSV_HEADERS,
            )
            self.assertEqual(members[2][0].value, 1)
            workbook.close()

    def test_sort_dialogs_by_group_title(self) -> None:
        class _Entity:
            def __init__(self, title: str) -> None:
                self.title = title

        class _Dialog:
            def __init__(self, title: str) -> None:
                self.entity = _Entity(title)

        dialogs = [_Dialog("Beta"), _Dialog("alpha"), _Dialog("Alpha")]
        _sort_dialogs_by_group_title(dialogs)
        self.assertEqual(
            [d.entity.title for d in dialogs],
            ["Alpha", "alpha", "Beta"],
        )

    def test_sort_members_by_display_name(self) -> None:
        rows = [
            {"display_name": "Bob"},
            {"display_name": "alice"},
            {"display_name": "Alice"},
        ]
        _sort_members_by_display_name(rows)
        self.assertEqual(
            [r["display_name"] for r in rows],
            ["Alice", "alice", "Bob"],
        )


if __name__ == "__main__":
    unittest.main()
